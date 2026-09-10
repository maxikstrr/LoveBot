#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
LoveBot — Brand-Kit Generator (Firmenprofil .docx · Präsentation .pptx · Fact-Sheet .rtf)
Erzeugt die fertigen Vorlagen unter Dokumente/BrandKit/.
Wiederverwendbar:  python3 scripts/make-brandkit.py   (aus dem LoveBot-Ordner)
"""
import json, os, shutil, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
OUT = os.path.join('Dokumente', 'BrandKit')
os.makedirs(OUT, exist_ok=True)

LOGO_PNG = os.path.join('public', 'assets', 'img', 'lovebot-logo.png')
LOGO_SVG = os.path.join('public', 'assets', 'img', 'lovebot-logo.svg')

# ---------- Farbwelt ----------------------------------------------------------
PURPLE = (0x7C, 0x3A, 0xED)   # 7C3AED  LoveBot-Violett
PINK   = (0xE9, 0x1E, 0x8C)   # E91E8C  Liebes-Pink
DARK   = (0x1B, 0x12, 0x33)   # Midnight
GREY   = (0x6B, 0x6B, 0x7A)
LIGHT  = (0xF7, 0xF3, 0xFC)
PINKHEX = 'E91E8C'
PURHEX  = '7C3AED'

# ---------- Live-Kennzahlen (Stand: Generierung) -------------------------------
def load_num():
    out = {}
    try:
        d = json.load(open('Database/Database.json'))
        out['users'] = len(d.get('users', {}))
        out['registered'] = sum(1 for u in d.get('users', {}).values()
                                if (u.get('registration') or {}).get('registered') is True)
        out['groups'] = len(d.get('groups', {}))
        out['active_groups'] = sum(1 for g in d.get('groups', {}).values()
                                   if g and g.get('active') is not False)
        out['bans'] = len(d.get('bans', {}))
    except Exception:
        pass
    try:
        s = json.load(open('Database/sessions.json'))
        out['sessions'] = len(s.get('sessions', {}))
        out['connected'] = sum(1 for x in s.get('sessions', {}).values()
                               if x.get('status') == 'CONNECTED')
    except Exception:
        pass
    try:
        a = json.load(open('Database/accounts.json'))
        out['accounts'] = len(a.get('accounts', {}))
    except Exception:
        pass
    try:
        lp = json.load(open('Database/loveplus.json'))
        out['couples'] = len(lp.get('couples', {}))
        out['pets'] = sum(1 for u in lp.get('users', {}).values() if u.get('pet'))
    except Exception:
        pass
    return out

N = load_num()
STAMP = datetime.date.today().strftime('%d.%m.%Y')
YEAR = datetime.date.today().strftime('%Y')

WEB_LINKS = [
    ('Website', 'https://maxichen.de'),
    ('LoveBot-Dashboard', 'https://maxichen.gamebot.me'),
    ('LoveBot-Kanal', 'https://whatsapp.com/channel/0029Vb8EH4IBqbrAu9LxUH3X'),
    ('GitHub', 'https://github.com/maxikstrr'),
    ('Instagram', 'https://www.instagram.com/max_.kstr'),
    ('TikTok', 'https://www.tiktok.com/@maxichensworld'),
]

# ==============================================================================
# 1) FIRMENPROFIL  (.docx)
# ==============================================================================
def build_docx():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # Grundschrift
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    st.font.color.rgb = RGBColor(*DARK)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')

    for s in doc.sections:
        s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.5)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

    def para(text='', size=11, bold=False, italic=False, color=DARK, align=None, space_after=6, space_before=0):
        p = doc.add_paragraph()
        if align is not None: p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        r = p.add_run(text)
        r.font.name = 'Calibri'; r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = RGBColor(*color)
        return p

    def rich(p, text, size=11, bold=False, color=DARK):
        r = p.add_run(text)
        r.font.name = 'Calibri'; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = RGBColor(*color)
        return r

    def heading(text, color=PURPLE, size=16, space_before=14):
        return para(text, size=size, bold=True, color=color, space_after=4, space_before=space_before)

    def bullets(items, color=DARK, size=11):
        for it in items:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(it)
            r.font.name = 'Calibri'; r.font.size = Pt(size); r.font.color.rgb = RGBColor(*color)

    def hr(color=PURHEX, size=14):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), str(size))
        bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), color)
        pBdr.append(bottom); pPr.append(pBdr)

    # ---- Kopfbereich: Logo steht IMMER oben rechts (Kopfzeile, jede Seite) ----
    hdr = doc.sections[0].header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.space_before = Pt(0)
    run = hp.add_run()
    run.add_picture(LOGO_PNG, height=Cm(1.5))
    doc.sections[0].header_distance = Cm(0.7)

    para('LOVEBOT', size=30, bold=True, color=PURPLE, space_after=0)
    p = para('', space_after=10)
    rich(p, 'Company Profile', size=15, bold=True, color=PINK)
    rich(p, '   ·   LoveBot by Maxichen  ·  ' + YEAR, size=12, color=GREY)
    hr()

    # ---- Über uns ----
    heading('Über LoveBot')
    para('LoveBot ist ein moderner WhatsApp-Begleit-Bot mit eigenem Web-Control-Panel '
         '(„Midnight Control“). LoveBot verbindet Community-Features, Unterhaltung, '
         'Moderation und Verwaltung in einem System — entwickelt und betrieben von Maxichen.',
         space_after=8)
    para('LoveBot ist nicht nur ein Bot: Über das Control-Panel lassen sich Gruppen, '
         'Nutzer, Rechte, Bot-Sessions und das gesamte Love-System live verwalten — '
         'sicher, mit Rollen & Rechten (Owner · Deputy · Admin · Supporter) und '
         'vollständigem Audit-Log.', space_after=8)

    # ---- Was LoveBot kann ----
    heading('Was LoveBot kann')
    bullets([
        '💜 Community & Liebe — Profile ($me), Paare, Herzen, Ehe & Beziehungs-Tools',
        '💰 Wallet & Ökonomie — Kupfer/Silber/Gold/Platin, tägliche Belohnungen, Bank',
        '🎮 Games & Spaß — Minispiele, Achievements, Haustiere, Level & Prestige',
        '🛡️ Moderation — Auto-Welcome/Goodbye, Kick-Meldungen, Promote/Demote',
        '🔗 Gruppen-Werkzeuge — eigene Gruppen ($sss), Badword-Filter, Anti-Link, $join/$leave',
        '🧰 Alltags-Tools — $afk, $ping, $cmd, Wetter, Musik & Medien-Befehle',
        '📢 Channel-Spiegel — Inhalte automatisch in Gruppen & Chats weiterleiten',
        '🖥️ Midnight Control — Web-Panel für Sessions, Nutzer, Gruppen, Logs & Rechte'
    ])

    # ---- Zahlen & Fakten ----
    heading('Zahlen & Fakten')
    para('Aktuelle Live-Kennzahlen (Stand ' + STAMP + '):', size=11, italic=True, color=GREY, space_after=6)
    rows = [
        ('💜 Nutzer in der Datenbank', str(N.get('users', 0))),
        ('✅ Registrierte Nutzer', str(N.get('registered', 0))),
        ('👥 Gruppen', str(N.get('groups', 0)) + '  ·  ' + str(N.get('active_groups', 0)) + ' aktiv'),
        ('🤖 Bot-Sessions', str(N.get('sessions', 0)) + '  ·  ' + str(N.get('connected', 0)) + ' verbunden'),
        ('🛡️ Dashboard-Accounts', str(N.get('accounts', 0))),
        ('💍 Liebespaare', str(N.get('couples', 0))),
        ('🚫 Verwaltete Sperren', str(N.get('bans', 0))),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = 'Light Shading Accent 1'
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = k; c1.text = v
        for cell, rr in ((c0, 0), (c1, 0)):
            for pgh in cell.paragraphs:
                for run_ in pgh.runs:
                    run_.font.name = 'Calibri'; run_.font.size = Pt(11)
                    if rr == 0: run_.font.bold = True
                    run_.font.color.rgb = RGBColor(*DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- Werte ----
    heading('Werte & Versprechen')
    bullets([
        '💜 Liebe im Detail — jedes Feature mit Herz, von Community-Mitgliedern für Community-Mitglieder.',
        '🛡️ Sicherheit zuerst — Passwörter werden nie gespeichert, Rechte & Audit-Log überall.',
        '🌙 Midnight vibes — eine Marke, die nachts am hellsten leuchtet.',
    ])

    # ---- Kontakt ----
    heading('Kontakt & Kanäle')
    for label, url in WEB_LINKS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        rich(p, '•  ' + label + ':  ', size=11, bold=True, color=PURPLE)
        rich(p, url, size=11, color=DARK)
    para('', space_after=2)
    para('🔗 maxichen.gamebot.me  ·  maxichen.de', size=10, color=GREY, space_after=0)

    # Fußzeile
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run('© ' + YEAR + ' LoveBot by Maxichen  ·  💜  ·  Alle Rechte vorbehalten.')
    fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(*GREY)

    path = os.path.join(OUT, 'LoveBot-Firmenprofil-' + YEAR + '.docx')
    doc.save(path)
    print('DOCX ->', path)

# ==============================================================================
# 2) PRÄSENTATION  (.pptx)
# ==============================================================================
def build_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn

    SW, SH = Inches(13.333), Inches(7.5)
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    C_PUR = RGBColor(0x7C, 0x3A, 0xED)
    C_PINK = RGBColor(0xE9, 0x1E, 0x8C)
    C_DARK = RGBColor(0x1B, 0x12, 0x33)
    C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    C_GREY = RGBColor(0xB9, 0xB2, 0xCD)

    def new_slide(bg=C_DARK):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        return s

    def txt(slide, l, t, w, h, text, size=18, bold=False, italic=False, color=C_WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(l, t, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
        lines = text.split('\n')
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run(); r.text = ln
            r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = 'Calibri'
        return tb

    def box(slide, l, t, w, h, fill, line=None, radius=False):
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            l, t, w, h)
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line; shp.line.width = Pt(1)
        return shp

    def accent_bar(slide, l, t, w, h, color=C_PINK):
        return box(slide, l, t, w, h, color)

    def logo(slide, right=True, width=Inches(1.05), top=Inches(0.42)):
        from pptx.util import Emu as E
        l = SW - width - Inches(0.5) if right else Inches(0.5)
        slide.shapes.add_picture(LOGO_PNG, int(l), int(top), height=int(width))

    def footer(slide, n):
        txt(slide, Inches(0.5), SH - Inches(0.42), Inches(8), Inches(0.3),
            'LoveBot by Maxichen  ·  maxichen.gamebot.me', size=10, color=C_GREY)
        txt(slide, SW - Inches(1.2), SH - Inches(0.46), Inches(0.7), Inches(0.3),
            str(n), size=12, color=C_GREY, align=PP_ALIGN.RIGHT)

    # --- Slide 1: Titel ---
    s = new_slide()
    accent_bar(s, 0, Inches(0.14), SW, Inches(0.09), C_PINK)
    accent_bar(s, 0, Inches(0.23), SW, Inches(0.045), C_PUR)
    logo(s, right=True, width=Inches(1.5), top=Inches(0.75))
    txt(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.4), 'LOVEBOT', size=72, bold=True, color=C_WHITE)
    txt(s, Inches(1.0), Inches(3.85), Inches(11.0), Inches(0.8), 'Das LoveBot-Control-Panel · Firmenvorstellung', size=26, color=C_PINK)
    txt(s, Inches(1.0), Inches(4.75), Inches(11.0), Inches(0.6), 'by Maxichen  ·  ' + YEAR, size=18, color=C_GREY)
    accent_bar(s, Inches(1.0), Inches(5.7), Inches(2.4), Inches(0.07), C_PUR)
    txt(s, Inches(0.5), SH - Inches(0.45), Inches(8), Inches(0.3), 'WhatsApp-Begleit-Bot & Midnight-Control-Panel', size=11, color=C_GREY)

    # --- Slide 2: Was ist LoveBot? ---
    s = new_slide()
    logo(s)
    txt(s, Inches(0.7), Inches(0.55), Inches(9), Inches(0.9), 'Was ist LoveBot?', size=36, bold=True, color=C_WHITE)
    accent_bar(s, Inches(0.72), Inches(1.45), Inches(1.6), Inches(0.07), C_PINK)
    body = [
        'LoveBot ist ein WhatsApp-Begleit-Bot für Communitys — mit Herz, Spiel und Ordnung.',
        'Er begleitet Gruppen täglich: Profile & Liebe, Wallet & Minispiele, Moderation & Tools.',
        'Dazu kommt das Midnight Control-Panel: Gruppen, Nutzer, Rechte, Bot-Sessions & Logs an einem Ort.',
        'Eine Marke, die nachts am hellsten leuchtet — entwickelt von Maxichen.',
    ]
    y = 1.9
    for b in body:
        txt(s, Inches(0.85), Inches(y), Inches(11.6), Inches(0.9), '•  ' + b, size=17, color=C_WHITE)
        y += 0.92
    footer(s, 2)

    # --- Slide 3: Module (Karten) ---
    s = new_slide()
    logo(s)
    txt(s, Inches(0.7), Inches(0.5), Inches(9), Inches(0.9), 'Module & Highlights', size=34, bold=True, color=C_WHITE)
    accent_bar(s, Inches(0.72), Inches(1.38), Inches(1.6), Inches(0.07), C_PINK)
    cards = [
        ('💜 Love & Community', 'Profile, Paare, Herzen & Ehe – das Herz von LoveBot.'),
        ('💰 Wallet & Economy', 'Kupfer bis Platin, tägliche Belohnungen, Bank & Level.'),
        ('🎮 Games & Achievements', 'Minispiele, Erfolge, Haustiere, Level & Prestige.'),
        ('🛡️ Moderation', 'Welcome/Goodbye, Kick-Meldung, Promote/Demote, Bans.'),
        ('🔗 Gruppen-Werkzeuge', 'Anti-Link, Badwords, $sss-Gruppen, $join & $leave.'),
        ('🖥️ Midnight Control', 'Sessions, Nutzer, Rechte & Logs im Web-Panel.'),
    ]
    cw, ch, gx, gy = Inches(3.86), Inches(2.15), Inches(0.25), Inches(0.3)
    x0, y0 = Inches(0.75), Inches(1.85)
    for i, (t, d) in enumerate(cards):
        cx = x0 + (i % 3) * (cw + gx)
        cy = y0 + (i // 3) * (ch + gy)
        card = box(s, cx, cy, cw, ch, RGBColor(0x2A, 0x1E, 0x4E), radius=True)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.2); tf.margin_top = Inches(0.2)
        p1 = tf.paragraphs[0]; r1 = p1.add_run(); r1.text = t
        r1.font.size = Pt(18); r1.font.bold = True; r1.font.color.rgb = C_PINK; r1.font.name = 'Calibri'
        p2 = tf.add_paragraph(); p2.space_before = Pt(10)
        r2 = p2.add_run(); r2.text = d
        r2.font.size = Pt(13); r2.font.color.rgb = C_WHITE; r2.font.name = 'Calibri'
    footer(s, 3)

    # --- Slide 4: Zahlen & Fakten ---
    s = new_slide()
    logo(s)
    txt(s, Inches(0.7), Inches(0.5), Inches(9), Inches(0.9), 'Zahlen & Fakten', size=34, bold=True, color=C_WHITE)
    accent_bar(s, Inches(0.72), Inches(1.38), Inches(1.6), Inches(0.07), C_PINK)
    txt(s, Inches(0.75), Inches(1.75), Inches(11), Inches(0.5),
        'Live-Kennzahlen · Stand ' + STAMP, size=15, italic=True, color=C_GREY)
    stats = [
        ('💜', str(N.get('users', 0)), 'Nutzer in der Datenbank'),
        ('👥', str(N.get('groups', 0)), 'Gruppen (' + str(N.get('active_groups', 0)) + ' aktiv)'),
        ('🤖', str(N.get('sessions', 0)), 'Bot-Sessions'),
        ('🛡️', str(N.get('accounts', 0)), 'Dashboard-Accounts'),
        ('💍', str(N.get('couples', 0)), 'Liebespaare'),
        ('✅', str(N.get('registered', 0)), 'registrierte Nutzer'),
    ]
    cw2, ch2, g2 = Inches(3.86), Inches(1.7), Inches(0.25)
    for i, (ic, num, lab) in enumerate(stats):
        cx = x0 = Inches(0.75) + (i % 3) * (cw2 + g2)
        cy = Inches(2.45) + (i // 3) * (ch2 + Inches(0.35))
        b = box(s, cx, cy, cw2, ch2, RGBColor(0x2A, 0x1E, 0x4E), radius=True)
        tf = b.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.2); tf.margin_top = Inches(0.22)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
        r = p.add_run(); r.text = ic + '  ' + num
        r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = C_PINK; r.font.name = 'Calibri'
        p2 = tf.add_paragraph(); p2.space_before = Pt(6)
        r2 = p2.add_run(); r2.text = lab
        r2.font.size = Pt(13); r2.font.color.rgb = C_WHITE; r2.font.name = 'Calibri'
    footer(s, 4)

    # --- Slide 5: Sicherheit / Werte ---
    s = new_slide()
    logo(s)
    txt(s, Inches(0.7), Inches(0.5), Inches(9), Inches(0.9), 'Sicherheit & Werte', size=34, bold=True, color=C_WHITE)
    accent_bar(s, Inches(0.72), Inches(1.38), Inches(1.6), Inches(0.07), C_PINK)
    vals = [
        ('🔐', 'Rollen & Rechte', 'Owner, Deputy, Admin & Supporter — alles granular steuerbar.'),
        ('📜', 'Audit-Log', 'Jede kritische Aktion wird protokolliert & nachvollziehbar.'),
        ('💜', 'Privacy first', 'Kein Standort-Tracking, DSGVO-Zustimmung integriert.'),
        ('🌙', 'Midnight Design', 'Ein Control-Panel, das nachts am schönsten ist.'),
    ]
    y = 1.95
    for ic, t, d in vals:
        b = box(s, Inches(0.85), Inches(y), Inches(11.6), Inches(1.0), RGBColor(0x26, 0x1A, 0x45), radius=True)
        tf = b.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(0.28); tf.margin_top = Inches(0.14)
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = ic + '  ' + t + '  —  ' + d
        r.font.size = Pt(16); r.font.color.rgb = C_WHITE; r.font.name = 'Calibri'
        r2 = p.add_run()
        y += 1.15
    footer(s, 5)

    # --- Slide 6: Kontakt ---
    s = new_slide()
    logo(s, width=Inches(1.3), top=Inches(1.1))
    txt(s, Inches(0.7), Inches(3.2), Inches(12), Inches(1.0), 'LoveBot by Maxichen', size=40, bold=True, color=C_WHITE)
    txt(s, Inches(0.7), Inches(4.15), Inches(12), Inches(0.6), 'Danke fürs Reinschauen — Liebe Grüße, Maxichen 💜', size=20, color=C_PINK)
    txt(s, Inches(0.7), Inches(5.35), Inches(12), Inches(0.5), 'maxichen.gamebot.me   ·   maxichen.de', size=16, color=C_GREY)
    accent_bar(s, Inches(0.7), Inches(6.6), Inches(3.0), Inches(0.06), C_PUR)
    txt(s, Inches(0.5), SH - Inches(0.45), Inches(8), Inches(0.3), 'LoveBot ☾ · Midnight Control', size=11, color=C_GREY)

    path = os.path.join(OUT, 'LoveBot-Praesentation-' + YEAR + '.pptx')
    prs.save(path)
    print('PPTX ->', path)

# ==============================================================================
# 3) FACT-SHEET (.rtf)
# ==============================================================================
def esc_rtf(t):
    t = t.replace('\\', '\\\\').replace('{', '\\{').replace('}', '\\}')
    out = []
    for ch in t:
        o = ord(ch)
        if o < 128:
            out.append(ch)
        else:
            out.append('\\u%d?' % o)
    return ''.join(out)

def build_rtf():
    L = []
    L.append(r'{\rtf1\ansi\deff0{\fonttbl{\f0 Calibri;}}')
    L.append(r'{\colortbl ;\red124\green58\blue237;\red233\green30\blue140;\red27\green18\blue51;\red107\green107\blue122;}')
    L.append(r'\paperw11906\paperh16838\margl1134\margr1134\margt850\margb850\f0\fs22 ')
    def sec(title):
        L.append(r'\pard\qc\f0\fs20\b\cf1 ' + esc_rtf(title) + r'\par\fs14\cf4\i ' + esc_rtf('LOVEBOT by Maxichen · ' + YEAR) + r'\i0\par\pard\fs2\par')
    # Marken-Zeile oben rechts
    L.append(r'\pard\qr\f0\fs18\cf2\b LOVE BOT ☾ \cf4\f0\b0\fs12 by Maxichen · ' + YEAR + r'\par\pard\fs2\par')
    L.append(r'\pard\qc\f0\fs44\b\cf1 LOVE\cf2 BOT\f0\fs18\par\pard\fs22 ')
    L.append(r'\pard\qc\fs18\cf4\i Fact-Sheet & Überblick\i0\par\pard\fs2\par')
    # Absatzhelfer
    def h(text):
        L.append(r'\pard\fs24\b\cf1\sa120 ' + esc_rtf(text) + r'\b0\par')
    def p(text, cf=3):
        L.append(r'\pard\fs22\cf' + str(cf) + r'\sa80\li280 ' + esc_rtf(text) + r'\par')
    def bullet(text):
        L.append(r'\pard\fs22\cf3\sa60\li500 ' + esc_rtf('•  ' + text) + r'\par')
    def kv(k, v):
        L.append(r'\pard\fs22\sa40\li280\cf3\b ' + esc_rtf(k) + r'\b0\tab ' + esc_rtf(v) + r'\par')

    h('Über LoveBot')
    p('LoveBot ist ein moderner WhatsApp-Begleit-Bot mit eigenem Web-Control-Panel (Midnight Control). Community, Unterhaltung, Moderation und Verwaltung in einem System — entwickelt und betrieben von Maxichen.')

    h('Highlights')
    bullet('Community & Liebe: Profile, Paare, Herzen & Ehe')
    bullet('Wallet & Economy: Kupfer, Silber, Gold, Platin, tägliche Belohnungen, Bank')
    bullet('Games: Minispiele, Achievements, Haustiere, Level & Prestige')
    bullet('Moderation: Auto-Welcome/Goodbye, Kick-Meldungen, Promote/Demote, Bans')
    bullet('Gruppen-Werkzeuge: Anti-Link, Badwords, $sss-Gruppen, $join/$leave')
    bullet('Midnight Control: Web-Panel für Sessions, Nutzer, Rechte & Logs')

    h('Zahlen & Fakten  (' + STAMP + ')')
    kv('Nutzer in der Datenbank', str(N.get('users', 0)))
    kv('Registrierte Nutzer', str(N.get('registered', 0)))
    kv('Gruppen', str(N.get('groups', 0)) + '  (' + str(N.get('active_groups', 0)) + ' aktiv)')
    kv('Bot-Sessions', str(N.get('sessions', 0)) + '  (' + str(N.get('connected', 0)) + ' verbunden)')
    kv('Dashboard-Accounts', str(N.get('accounts', 0)))
    kv('Liebespaare', str(N.get('couples', 0)))
    kv('Verwaltete Sperren', str(N.get('bans', 0)))

    h('Kontakt & Kanäle')
    for label, url in WEB_LINKS:
        kv(label, url)

    L.append(r'\par\pard\sa60\fs16\cf4\qc ' + esc_rtf('🔗 maxichen.gamebot.me  ·  maxichen.de') + r'\par')
    L.append(r'\pard\sa60\fs16\cf4\qc ' + esc_rtf('© ' + YEAR + ' LoveBot by Maxichen · Alle Rechte vorbehalten.') + r'\par')
    L.append('}')
    path = os.path.join(OUT, 'LoveBot-Factsheet-' + YEAR + '.rtf')
    with open(path, 'w', encoding='utf-8', errors='replace') as f:
        f.write(''.join(L))
    print('RTF  ->', path)

# ==============================================================================
# 4) Kopien + README
# ==============================================================================
def copies():
    shutil.copy2(LOGO_PNG, os.path.join(OUT, 'LoveBot-Logo.png'))
    shutil.copy2(LOGO_SVG, os.path.join(OUT, 'LoveBot-Logo.svg'))
    lines = [
        'LOVEBOT ☾ — Brand-Kit ' + YEAR,
        '=' * 42,
        '',
        'Inhalt dieses Ordners:',
        '  • LoveBot-Firmenprofil-' + YEAR + '.docx    – Unternehmensprofil (Word)',
        '  • LoveBot-Praesentation-' + YEAR + '.pptx  – Vorstellung (PowerPoint)',
        '  • LoveBot-Factsheet-' + YEAR + '.rtf       – Fact-Sheet (RTF / WordPad)',
        '  • LoveBot-Logo.png                        – Logo (transparent, 1024 px)',
        '  • LoveBot-Logo.svg                        – Logo (Vektor, skalierbar)',
        '',
        'Kennzahlen in den Dateien: Stand ' + STAMP + '.',
        'Neu erzeugen (nur für die Entwicklung): python3 scripts/make-brandkit.py',
        '',
        'Marke: LoveBot by Maxichen  ·  maxichen.gamebot.me  ·  maxichen.de',
    ]
    with open(os.path.join(OUT, 'LIESMICH-Brand-Kit.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print('README + Logos ->', OUT)

if __name__ == '__main__':
    build_docx()
    build_pptx()
    build_rtf()
    copies()
    print('Brand-Kit fertig ✓')
